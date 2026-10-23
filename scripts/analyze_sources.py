"""
Обзор исходных таблиц: листы, размеры, формулы и где лежат персональные данные.

Запуск: .venv-sheets/Scripts/python.exe scripts/analyze_sources.py
"""

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ATTACHMENTS = Path(r"C:\Users\nikyu\AppData\Local\hermes\attachments")

FILES = [
    "Копия  Планы продаж НР.xlsx",
    "г. Новосибирск, пр. Карла Маркса, д.7 (1).xlsx",
    "г. Новосибирск, пр. Карла Маркса, д.7.xlsx",
    "График НР 2026.xlsx",
    "Комбо НСК .xlsx",
    "Единый Калькулятор Combo .xlsx",
    "НР_ Обновленная мотивация Т2 15.10.2026.xlsx",
]

# Признаки персональных данных.
PHONE_RE = re.compile(r"(\+7|8)[\s\-(]?\d{3}[\s\-)]?\d{3}[\s\-]?\d{2}[\s\-]?\d{2}")
EMAIL_RE = re.compile(r"[\w.\-]+@[\w.\-]+\.\w+")
# ФИО: три слова с заглавных букв, либо Фамилия И.О.
FIO_RE = re.compile(r"^[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+(\s+[А-ЯЁ][а-яё]+)?$")
TABULAR_FIO_RE = re.compile(
    r"^[А-ЯЁ][а-яё\-]+\s+[А-ЯЁ]\.\s?[А-ЯЁ]\.?\s*$|^[А-ЯЁ][а-яё\-]+\s+[А-ЯЁ][а-яё\-]+$"
)

KEYWORD_COLUMNS = (
    "фио",
    "ф.и.о",
    "ф и о",
    "сотрудник",
    "менеджер",
    "продавец",
    "табел",
    "телефон",
    "телеф",
    "почта",
    "email",
    "логин",
    "login",
    "адрес",
    "должност",
)


def looks_personal(value) -> bool:
    if not isinstance(value, str):
        return False
    text = value.strip()
    if not text or text.startswith("="):
        return False
    if PHONE_RE.search(text) or EMAIL_RE.search(text):
        return True
    if FIO_RE.match(text) or TABULAR_FIO_RE.match(text):
        return True
    return False


def analyze(path: Path) -> dict:
    print(f"\n{'=' * 78}\n{path.name}\n{'=' * 78}")
    workbook = load_workbook(path, data_only=False, read_only=False)

    report = {"file": path.name, "sheets": []}

    for sheet in workbook.worksheets:
        formulas = 0
        personal_samples = []
        keyword_headers = []

        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    formulas += 1
                    continue
                if isinstance(value, str):
                    lowered = value.strip().lower()
                    if lowered in KEYWORD_COLUMNS or (
                        len(lowered) < 30
                        and any(keyword in lowered for keyword in KEYWORD_COLUMNS)
                    ):
                        keyword_headers.append((cell.coordinate, value.strip()))
                        continue
                if looks_personal(value):
                    if len(personal_samples) < 8:
                        personal_samples.append((cell.coordinate, value.strip()))

        info = {
            "title": sheet.title,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "formulas": formulas,
            "keyword_headers": keyword_headers[:12],
            "personal_samples": personal_samples,
        }
        report["sheets"].append(info)

        print(f"\n  лист «{sheet.title}» — {sheet.max_row}×{sheet.max_column}, формул: {formulas}")
        if keyword_headers:
            print(f"    колонки с персональными данными: {keyword_headers[:8]}")
        if personal_samples:
            print(f"    примеры персональных значений: {personal_samples[:5]}")
        if formulas == 0 and sheet.max_row > 20:
            print("    (!) формул нет — возможно, это сводные значения")

    workbook.close()
    return report


def main() -> int:
    reports = []
    for name in FILES:
        path = ATTACHMENTS / name
        if not path.exists():
            print(f"НЕТ ФАЙЛА: {name}")
            continue
        try:
            reports.append(analyze(path))
        except Exception as error:  # noqa: BLE001
            print(f"ОШИБКА на {name}: {error}")

    out = Path("docs/data/source_inventory.json")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(reports, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\n\nОтчёт сохранён: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
