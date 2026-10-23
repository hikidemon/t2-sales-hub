"""
Где именно лежат ФИО в «Карла Маркса» и «График НР 2026»:
находим колонки с заголовками-подсказками и сами значения.

Запуск: .venv-sheets/Scripts/python.exe scripts/inspect_names.py
"""

import sys
from pathlib import Path

from openpyxl import load_workbook

ATTACHMENTS = Path(r"C:\Users\nikyu\AppData\Local\hermes\attachments")

PERSON_HEADERS = (
    "фио",
    "ф.и.о",
    "ф и о",
    "фамилия",
    "сотрудник",
    "продавец",
    "менеджер",
    "табел",
    "телефон",
    "почта",
    "email",
    "логин",
)


def is_person_header(text) -> bool:
    if not isinstance(text, str):
        return False
    lowered = text.strip().lower()
    if not lowered or len(lowered) > 40:
        return False
    return any(keyword in lowered for keyword in PERSON_HEADERS)


def report(path: Path, max_sheets: int = 4) -> None:
    print(f"\n{'=' * 78}\n{path.name}\n{'=' * 78}")

    workbook = load_workbook(path, data_only=False, read_only=False)
    for sheet in workbook.worksheets[:max_sheets]:
        print(f"\n--- лист «{sheet.title}» ({sheet.max_row}×{sheet.max_column}) ---")

        # Ищем колонки, чей заголовок говорит о персональных данных.
        person_columns = {}
        for row in sheet.iter_rows(min_row=1, max_row=min(12, sheet.max_row)):
            for cell in row:
                if is_person_header(cell.value):
                    person_columns.setdefault(cell.column, []).append(
                        (cell.row, str(cell.value).strip())
                    )

        if person_columns:
            print(f"  колонки с персональными заголовками: {sorted(person_columns)}")
            for column, headers in sorted(person_columns.items())[:4]:
                print(f"    колонка {column}: заголовки {headers[:3]}")
                for row_index in range(1, min(sheet.max_row, 12) + 1):
                    value = sheet.cell(row=row_index, column=column).value
                    if value not in (None, ""):
                        print(f"      {sheet.cell(row=row_index, column=column).coordinate} = {value!r}")

        # Все строки, где встречаются «человеческие» значения, без заголовка.
        if not person_columns:
            interesting = []
            for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 40)):
                for cell in row:
                    value = cell.value
                    if isinstance(value, str) and not value.startswith("="):
                        text = value.strip()
                        if "\n" in text or text.startswith("Ст "):
                            interesting.append((cell.coordinate, text.replace("\n", " / ")))
            for coordinate, text in interesting[:14]:
                print(f"    {coordinate} = {text!r}")

    workbook.close()


def main() -> int:
    for name in (
        "г. Новосибирск, пр. Карла Маркса, д.7 (1).xlsx",
        "График НР 2026.xlsx",
    ):
        report(ATTACHMENTS / name)
    return 0


if __name__ == "__main__":
    sys.exit(main())
