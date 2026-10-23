"""
Сверка: какие листы-ФИО переименованы, а какие остались.

Запуск: .venv-sheets/Scripts/python.exe scripts/check_renames.py
"""

import sys
from pathlib import Path

from openpyxl import load_workbook

ATTACHMENTS = Path(r"C:\Users\nikyu\AppData\Local\hermes\attachments")
SOURCE = ATTACHMENTS / "Копия  Планы продаж НР.xlsx"
REPORT = Path("build/anonymized/anonymization-report.json")


def main() -> int:
    import json

    report = json.loads(REPORT.read_text(encoding="utf-8"))
    renamed = {}
    for entry in report["files"]:
        if entry["file"].startswith("Копия"):
            renamed = entry["renamed_sheets"]

    workbook = load_workbook(SOURCE, read_only=True)
    titles = workbook.sheetnames
    workbook.close()

    missed = [title for title in titles if title not in renamed]
    done = [title for title in titles if title in renamed]

    print(f"всего листов: {len(titles)}")
    print(f"переименовано: {len(done)}")
    print(f"пропущено:     {len(missed)}")
    print()
    for title in missed[:40]:
        print(f"  ПРОПУЩЕН: {title}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
