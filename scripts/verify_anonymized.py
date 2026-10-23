"""
Проверка обезличенной книги «Карла Маркса»: что стало в графе фамилий
и в формулах, которые на них ссылаются.

Запуск: .venv-sheets/Scripts/python.exe scripts/verify_anonymized.py
"""

import sys
from pathlib import Path

from openpyxl import load_workbook

SOURCE = Path(
    r"C:\Users\nikyu\AppData\Local\hermes\attachments"
    r"\г. Новосибирск, пр. Карла Маркса, д.7 (1).xlsx"
)
TARGET = Path("build/anonymized/г. Новосибирск, пр. Карла Маркса, д.7 (1).xlsx")


def main() -> int:
    if not TARGET.exists():
        print(f"Нет результата: {TARGET}")
        return 1

    source = load_workbook(SOURCE, data_only=False)
    target = load_workbook(TARGET, data_only=False)

    print("листы оригинала:", source.sheetnames)
    print("листы результата:", target.sheetnames)
    print()

    sheet_source = source["график"]
    sheet_target = target["график"]

    print("графа B листа «график» (Было → Стало):")
    replaced = 0
    total = 0
    for row in range(1, sheet_source.max_row + 1):
        before = sheet_source.cell(row=row, column=2).value
        after = sheet_target.cell(row=row, column=2).value
        if before is None and after is None:
            continue
        if row > 15 and row % 20 != 0:
            continue
        total += 1
        if before != after:
            replaced += 1
        print(f"  B{row}: {before!r} {'  → ' + repr(after) if before != after else ''}")

    print(f"\n  показано строк: {total}, из них заменено: {replaced}")

    print()
    print("формулы «Лист_продаж», ссылающиеся на «график» (первые 6):")
    sales_source = source["Лист_продаж"]
    sales_target = target["Лист_продаж"]
    shown = 0
    for row in sales_source.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, str) and value.startswith("=") and "график" in value:
                after = sales_target[cell.coordinate].value
                same = "то же" if after == value else "ИЗМЕНЕНО"
                print(f"  {cell.coordinate}: {value}  [{same}]")
                shown += 1
                if shown >= 6:
                    break
        if shown >= 6:
            break

    source.close()
    target.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
