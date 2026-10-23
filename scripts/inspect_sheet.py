"""
Детальный разбор одного листа: где лежат ФИО, как устроены формулы,
какие ссылки на листы используются.

Запуск: .venv-sheets/Scripts/python.exe scripts/inspect_sheet.py
"""

import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ATTACHMENTS = Path(r"C:\Users\nikyu\AppData\Local\hermes\attachments")

SHEET_REF_RE = re.compile(r"'([^']+)'!|([A-Za-zА-Яа-яЁё0-9_\- ]+)!")


def dump(path: Path, sheet_name: str, rows: int = 14) -> None:
    workbook = load_workbook(path, data_only=False)
    sheet = workbook[sheet_name]

    print(f"\n{'=' * 78}")
    print(f"{path.name} → лист «{sheet_name}» ({sheet.max_row}×{sheet.max_column})")
    print("=" * 78)

    print("\n--- первые строки ---")
    for row in sheet.iter_rows(min_row=1, max_row=min(rows, sheet.max_row)):
        filled = [
            f"{cell.coordinate}={cell.value!r}"
            for cell in row
            if cell.value not in (None, "")
        ]
        if filled:
            print("  " + " | ".join(filled[:9]))

    formulas = []
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formulas.append((cell.coordinate, cell.value))

    print(f"\n--- формул: {len(formulas)}, примеры ---")
    for coordinate, formula in formulas[:12]:
        print(f"  {coordinate}: {formula}")

    refs = set()
    for _, formula in formulas:
        for match in SHEET_REF_RE.finditer(formula):
            name = match.group(1) or match.group(2)
            if name:
                refs.add(name.strip())
    print(f"\n--- ссылки на листы в формулах ({len(refs)}) ---")
    for name in sorted(refs)[:30]:
        print(f"  «{name}»")

    workbook.close()


def main() -> int:
    plans = ATTACHMENTS / "Копия  Планы продаж НР.xlsx"
    store = ATTACHMENTS / "г. Новосибирск, пр. Карла Маркса, д.7 (1).xlsx"

    workbook = load_workbook(plans, read_only=True)
    first = workbook.sheetnames[0]
    workbook.close()
    dump(plans, first)

    workbook = load_workbook(store, read_only=True)
    first_store = workbook.sheetnames[0]
    workbook.close()
    dump(store, first_store, rows=20)

    return 0


if __name__ == "__main__":
    sys.exit(main())
