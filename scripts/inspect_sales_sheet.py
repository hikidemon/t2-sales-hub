"""
Как устроен «Лист_продаж»: где фамилии и что от них зависит.

Запуск: .venv-sheets/Scripts/python.exe scripts/inspect_sales_sheet.py
"""

import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ATTACHMENTS = Path(r"C:\Users\nikyu\AppData\Local\hermes\attachments")
TARGET = ATTACHMENTS / "г. Новосибирск, пр. Карла Маркса, д.7 (1).xlsx"

PERSON_HEADERS = (
    "фио",
    "ф.и.о",
    "фамилия",
    "сотрудник",
    "продавец",
    "менеджер",
    "табел",
)


def main() -> int:
    workbook = load_workbook(TARGET, data_only=False)
    print("листы:", workbook.sheetnames)

    sheet = workbook["Лист_продаж"]
    print(f"\n«Лист_продаж» {sheet.max_row}×{sheet.max_column}")

    # Персональные заголовки в первых строках.
    person_columns = {}
    for row in sheet.iter_rows(min_row=1, max_row=12):
        for cell in row:
            if isinstance(cell.value, str):
                lowered = cell.value.strip().lower()
                if lowered and len(lowered) < 40 and any(
                    keyword in lowered for keyword in PERSON_HEADERS
                ):
                    person_columns.setdefault(cell.column, []).append(
                        (cell.row, cell.value.strip())
                    )

    print("\nколонки с персональными заголовками:")
    for column, headers in sorted(person_columns.items()):
        print(f"  колонка {column}: {headers[:3]}")
        for row_index in range(1, 20):
            value = sheet.cell(row=row_index, column=column).value
            if value not in (None, ""):
                print(
                    f"    {sheet.cell(row=row_index, column=column).coordinate} = {value!r}"
                )

    # Формулы, которые ссылаются на график или импорт.
    print("\nформулы со ссылками на «график»/«импорт»/«Планы»:")
    count = 0
    for row in sheet.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, str) and value.startswith("="):
                if any(
                    token in value
                    for token in ("график", "импорт", "Планы", "ВП")
                ):
                    print(f"  {cell.coordinate}: {value[:220]}")
                    count += 1
                    if count >= 10:
                        break
        if count >= 10:
            break

    # Что лежит в первых строках по колонкам.
    print("\nпервая строка таблицы (по колонкам):")
    for column in range(1, min(sheet.max_column, 40) + 1):
        values = []
        for row_index in range(1, 8):
            value = sheet.cell(row=row_index, column=column).value
            if value not in (None, ""):
                values.append(f"{row_index}:{value!r}")
        if values:
            print(f"  колонка {column}: " + " | ".join(str(v)[:60] for v in values[:4]))

    workbook.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
